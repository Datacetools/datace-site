import csv
import gc
import re
import io
import logging
from datetime import date
from pathlib import Path
import pdfplumber
import pandas as pd
from pdfminer.converter import PDFLayoutAnalyzer
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
logger = logging.getLogger('F24Splitter')
PDFLayoutAnalyzer.paint_path = lambda self, gstate, stroke, fill, evenodd, path: None
FIRMA_TESTO = 'Datace · datacetools.com'
FIRMA_LINK = 'https://datacetools.com/?utm_source=excel&utm_medium=file'
FIRMA_LOGO = Path(__file__).parent / 'firma-datace.png'

def _firma_datace(ws):
    try:
        from openpyxl.drawing.image import Image
        from openpyxl.worksheet.cell_range import CellRange
        from openpyxl.utils.cell import coordinate_from_string
        logo = Image(str(FIRMA_LOGO))
        logo.width = logo.height = 28
        unite = [str(r) for r in ws.merged_cells.ranges]
        for r in unite:
            ws.unmerge_cells(r)
        altezze = {i: d.height for i, d in ws.row_dimensions.items() if d.height}
        ws.insert_rows(1)
        for r in unite:
            cr = CellRange(r)
            cr.shift(row_shift=1)
            ws.merge_cells(cr.coord)
        for i in sorted(altezze, reverse=True):
            ws.row_dimensions[i + 1].height = altezze[i]
            ws.row_dimensions[i].height = None
        if ws.freeze_panes:
            col, riga = coordinate_from_string(ws.freeze_panes)
            ws.freeze_panes = f'{col}{riga + 1}'
        ws.row_dimensions[1].height = 30
        cella = ws.cell(1, 1, FIRMA_TESTO)
        cella.hyperlink = FIRMA_LINK
        cella.font = Font(size=11, bold=True, color='FFFFFF', underline='single')
        cella.fill = PatternFill('solid', fgColor='0F2744')
        cella.alignment = Alignment(indent=6, vertical='center')
        if ws.max_column > 1:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ws.max_column)
        from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
        logo.anchor = TwoCellAnchor(editAs='twoCell', _from=AnchorMarker(col=0, colOff=45000, row=0, rowOff=45000), to=AnchorMarker(col=0, colOff=335000, row=0, rowOff=335000))
        ws.add_image(logo)
    except Exception as e:
        logger.warning('Banner non inserito: %s', type(e).__name__)
NON_DISPONIBILE = 'Codice tributo (descrizione non disponibile)'
with open(Path(__file__).parent / 'dati' / 'codici_tributo.csv', encoding='utf-8') as _f:
    TRIBUTI: dict = {r['codice']: (r['descrizione ufficiale'] or NON_DISPONIBILE, r['sezione']) for r in csv.DictReader(_f, delimiter=';')}
SEZIONE_COLORI = {'ERARIO': '#D6E4FF', 'REGIONI': '#D6F5D6', 'IMU': '#FFF3CD', 'INPS': '#FCE4EC', 'ALTRI ENTI': '#F3E5F5', '?': '#F5F5F5'}
_IMPORTO_RE = re.compile('^\\d{1,3}(?:\\.\\d{3})*,\\d{2}$')
_YEAR_RE = re.compile('\\b(20\\d{2})\\b')
_YEAR_MMRE = re.compile('^\\d{2}(20\\d{2})$')
_TRIB_RE = re.compile('(?<!\\d)(\\d{4})(?!\\d)')
_SIGN_TOKENS = {'-', '−', '–', '—'}
_PAREN_NEG_RE = re.compile('^\\((\\d{1,3}(?:\\.\\d{3})*,\\d{2})\\)$')

def _parse_importo(text):
    t = (text or '').strip()
    sign = 1
    m = _PAREN_NEG_RE.match(t)
    if m:
        sign, t = (-1, m.group(1))
    elif t[:1] in _SIGN_TOKENS:
        sign, t = (-1, t[1:].strip())
    if not _IMPORTO_RE.match(t):
        return None
    return sign * float(t.replace('.', '').replace(',', '.'))

def _line_signed_amounts(texts):
    out, pending = ([], 1)
    for t in texts:
        if t in _SIGN_TOKENS:
            pending = -1
            continue
        v = _parse_importo(t)
        if v is None:
            pending = 1
            continue
        if pending == -1 and v > 0:
            v = -v
        pending = 1
        out.append(v)
    return out
_BANCA_RE = re.compile('DELEGA IRREVOCABILE A[:\\s]+(.+?)(?=\\s*MODELLO|\\s*PER\\s+L)', re.IGNORECASE | re.DOTALL)
_BANCA_KEYWORDS = {'banca', 'credito', 'credit', 'poste', 'intesa', 'unicredit', 'bnl', 'monte', 'mps', 'cassa', 'banco', 'mediobanca', 'agricole', 'sanpaolo', 'fineco', 'ing', 'deutsche', 'ubi', 'bper', 'credem', 'carige'}

def _extract_banca(page_text: str) -> str:
    m = _BANCA_RE.search(page_text)
    if not m:
        return ''
    raw = m.group(1).strip().replace('\n', ' ').strip()
    if any((kw in raw.lower() for kw in _BANCA_KEYWORDS)):
        return raw
    return ''
_NUM = '\\d{1,3}(?:\\.\\d{3})*,\\d{2}'
_TOTALE_RE = re.compile(f'TOTALE\\s+[ACEGIM]\\s*({_NUM})?\\s*[BDFHLN]+\\s*({_NUM})?')
_SALDO_FINALE_RE = re.compile(f'EURO\\s*\\+\\s*({_NUM})')

def _totali_stampati(page_text):
    fin = _SALDO_FINALE_RE.search(page_text)
    if not fin:
        return None
    num = lambda s: float(s.replace('.', '').replace(',', '.')) if s else 0.0
    tot = list(_TOTALE_RE.finditer(page_text))
    return (round(sum((num(m.group(1)) for m in tot)), 2), round(sum((num(m.group(2)) for m in tot)), 2), num(fin.group(1)))
_DATA_TOKEN_RE = re.compile('^[\\d./\\-]+$')

def _extract_data_versamento(words):
    estremi = next((w for w in words if w['text'].upper() == 'ESTREMI'), None)
    if not estremi:
        return None
    sotto = lambda t: [w for w in words if w['text'].lower() == t and w['top'] > estremi['top']]
    giorno, anno = (sotto('giorno'), sotto('anno'))
    if not (giorno and anno):
        return None
    g = min(giorno, key=lambda w: w['top'])
    a = min(anno, key=lambda w: abs(w['top'] - g['top']))
    cifre = sorted((w for w in words if _DATA_TOKEN_RE.match(w['text']) and w.get('upright', True) and (g['x0'] - 25 <= w['x0']) and (w['x1'] <= a['x1'] + 25) and (estremi['bottom'] <= w['top'] <= g['bottom'] + 22)), key=lambda w: w['x0'])
    sep = '' if all((len(w['text']) == 1 for w in cifre)) else ' '
    n = re.findall('\\d+', sep.join((w['text'] for w in cifre)))
    if len(n) == 1 and len(n[0]) in (6, 8):
        n = [n[0][:2], n[0][2:4], n[0][4:]]
    if len(n) != 3:
        return None
    gg, mm, aa = (int(x) for x in n)
    if aa < 100:
        aa += 2000
    try:
        return date(aa, mm, gg).strftime('%d/%m/%Y')
    except ValueError:
        return None

def _find_section_boundaries(lines):
    labels = {'SEZIONE ERARIO': 'ERARIO', 'SEZIONE INPS': 'INPS', 'SEZIONE REGIONI': 'REGIONI', 'SEZIONE IMU': 'IMU', 'SEZIONE ALTRI ENTI': 'ALTRI ENTI', 'ALTRI ENTI PREVIDENZIALI': 'ALTRI ENTI', 'SEZIONE ERARIO ED ALTRO': 'ERARIO'}
    bounds = []
    for ln in lines:
        line_txt = ' '.join((w['text'] for w in ln)).upper()
        for key, sec in labels.items():
            if key in line_txt:
                bounds.append((ln[0]['top'], sec))
                break
    return sorted(bounds)

def _section_at_y(bounds, y, fallback='?'):
    sec = fallback
    for top, name in bounds:
        if y >= top:
            sec = name
    return sec

def _find_phrase_bbox(words, tokens):
    n = len(tokens)
    for i in range(len(words) - n + 1):
        seg = words[i:i + n]
        if [w['text'].lower() for w in seg] == [t.lower() for t in tokens]:
            return (min((w['x0'] for w in seg)), min((w['top'] for w in seg)), max((w['x1'] for w in seg)), max((w['bottom'] for w in seg)))
    return None

def _column_bands(words, page_width, page_height):
    trib_hdr = _find_phrase_bbox(words, ['codice', 'tributo']) or _find_phrase_bbox(words, ['cod.', 'tributo'])
    deb_hdr = _find_phrase_bbox(words, ['importi', 'a', 'debito', 'versati']) or _find_phrase_bbox(words, ['importi', 'a', 'debito']) or _find_phrase_bbox(words, ['a', 'debito'])
    cre_hdr = _find_phrase_bbox(words, ['importi', 'a', 'credito', 'compensati']) or _find_phrase_bbox(words, ['importi', 'a', 'credito']) or _find_phrase_bbox(words, ['a', 'credito'])
    if not (trib_hdr and deb_hdr):
        return None
    W = page_width
    mid = lambda b: (b[0] + b[2]) / 2
    if cre_hdr:
        xs = sorted([(mid(trib_hdr), 'trib'), (mid(deb_hdr), 'deb'), (mid(cre_hdr), 'cre')], key=lambda x: x[0])
    else:
        xs = sorted([(mid(trib_hdr), 'trib'), (mid(deb_hdr), 'deb')], key=lambda x: x[0])
    bands = {}
    for i, (x, label) in enumerate(xs):
        x_left = xs[i - 1][0] if i > 0 else 0
        x_right = xs[i + 1][0] if i < len(xs) - 1 else W
        bands[label] = (max(0, (x_left + x) / 2), 0, min(W, (x + x_right) / 2), page_height)
    return bands

def _in_bbox(w, bbox):
    return w['x0'] >= bbox[0] and w['x1'] <= bbox[2] and (w['top'] >= bbox[1]) and (w['bottom'] <= bbox[3])

def _group_lines(words_sorted, tolerance=2.0):
    lines, line, prev_top = ([], [], None)
    for w in words_sorted:
        top = round(w['top'], 1)
        if prev_top is None or abs(top - prev_top) <= tolerance:
            line.append(w)
            prev_top = top
        else:
            if line:
                lines.append(line)
            line = [w]
            prev_top = top
    if line:
        lines.append(line)
    return lines
_CAUSALE_INPS_RE = re.compile('^(?=.*[A-Z])[A-Z0-9]{2,4}$')

def _causale_inps(ln):
    return next((w['text'] for w in ln if _CAUSALE_INPS_RE.match(w['text'])), None)
_INAIL_CAMPI = ('\\d{5}', '\\d{1,8}', '\\d{2}', '\\d{6}', '[A-Z]')
with open(Path(__file__).parent / 'dati' / 'sedi_inail.csv', encoding='utf-8') as _f:
    SEDI_INAIL = {r[0]: r[1] for r in list(csv.reader(_f, delimiter=';'))[1:]}

def _riga_inail(ln):
    t = [w['text'] for w in sorted(ln, key=lambda w: w['x0'])]
    for i in range(len(t) - 4):
        if all((re.fullmatch(p, x) for p, x in zip(_INAIL_CAMPI, t[i:i + 5]))):
            return (t[i], t[i + 4])
    return None

def _descrizione_inail(sede, causale):
    nome = SEDI_INAIL.get(sede)
    return f'INAIL, sede {sede} {nome}, causale {causale}' if nome else f'INAIL, sede {sede}, causale {causale}'
_CODICI_COME_ANNI = {'2001', '2002', '2003'}

def _is_valid_code(code):
    if re.match('^20\\d{2}$', code):
        return code in _CODICI_COME_ANNI
    if code in TRIBUTI:
        return True
    if re.match('^00\\d{2}$', code):
        return False
    if int(code) < 100:
        return False
    return True

def parse_pdf(source, progress_cb=None, stats=None) -> list:
    rows = []
    if isinstance(source, (bytes, bytearray)):
        pdf_source = io.BytesIO(source)
        filename = 'documento.pdf'
        del source
    elif isinstance(source, io.IOBase):
        pdf_source = source
        filename = 'documento.pdf'
    else:
        pdf_source = Path(source)
        filename = Path(source).name
    with pdfplumber.open(pdf_source) as pdf:
        total = len(pdf.pages)
        text_pages = 0
        controllate, non_tornano = (0, [])
        logger.info('Inizio parsing (%d pagine)', total)
        for page_idx, page in enumerate(pdf.pages, start=1):
            if progress_cb:
                progress_cb(page_idx, total)
            page_text, inizio = ('', len(rows))
            try:
                page_text = page.extract_text() or ''
                if page_text.strip():
                    text_pages += 1
                banca = _extract_banca(page_text)
                anno_page = next(iter(_YEAR_RE.findall(page_text)), None)
                if not anno_page:
                    for tok in page_text.split():
                        m = _YEAR_MMRE.match(tok)
                        if m:
                            anno_page = m.group(1)
                            break
                words_sorted = sorted(page.extract_words(use_text_flow=True, keep_blank_chars=False), key=lambda w: (round(w['top'], 1), w['x0']))
                data_vers = _extract_data_versamento(words_sorted)
                bands = _column_bands(words_sorted, page.width, page.height)
                lines = _group_lines(words_sorted)
                sec_bounds = _find_section_boundaries(lines)
                del words_sorted
                for ln in lines:
                    if not any((_parse_importo(w['text']) is not None for w in ln)):
                        continue
                    y = ln[0]['top']
                    sec_ln = _section_at_y(sec_bounds, y, '?')
                    trib = None
                    causale = _causale_inps(ln) if sec_ln == 'INPS' else None
                    inail = _riga_inail(ln) if sec_ln == 'ALTRI ENTI' else None
                    if not (causale or inail):
                        for w in ln:
                            m = _TRIB_RE.match(w['text'])
                            if m and _is_valid_code(m.group(1)):
                                trib = m.group(1)
                                break
                        if not trib:
                            continue
                    if inail:
                        trib, sezione, desc = ('INAIL', 'ALTRI ENTI', _descrizione_inail(*inail))
                    elif causale:
                        trib, sezione = (causale, 'INPS')
                        desc = f'Contributi INPS – causale {causale}'
                    elif trib in TRIBUTI:
                        desc, sezione = TRIBUTI[trib]
                        if not sezione:
                            sezione = sec_ln if sec_ln != '?' else 'ERARIO'
                    else:
                        desc = f'Codice tributo {trib} (non classificato)'
                        sezione = sec_ln if sec_ln != '?' else 'ERARIO'
                    anno = next((m2.group(1) for w in ln for m2 in [_YEAR_RE.search(w['text'])] if m2), None)
                    if anno is None:
                        anno = next((m2.group(1) for w in ln for m2 in [_YEAR_MMRE.match(w['text'])] if m2), anno_page)
                    if anno is not None:
                        try:
                            anno_int = int(anno)
                            if not 2000 <= anno_int <= 2099:
                                anno = anno_page
                        except (ValueError, TypeError):
                            anno = anno_page
                    if inail:
                        anno = anno_page
                    if anno is None:
                        anno = 'N/A'
                        logger.debug('Anno non trovato per tributo %s pagina %d', trib, page_idx)
                    imp_deb = imp_cre = None
                    if bands:
                        pending_sign = 1
                        for w in ln:
                            txt = w['text']
                            if txt in _SIGN_TOKENS:
                                pending_sign = -1
                                continue
                            val = _parse_importo(txt)
                            if val is None:
                                pending_sign = 1
                                continue
                            if pending_sign == -1 and val > 0:
                                val = -val
                            pending_sign = 1
                            cx = (w['x0'] + w['x1']) / 2
                            if bands['deb'][0] <= cx <= bands['deb'][2]:
                                imp_deb = val
                            elif 'cre' in bands and bands['cre'][0] <= cx <= bands['cre'][2]:
                                imp_cre = val
                    else:
                        signed = _line_signed_amounts([w['text'] for w in ln])
                        if signed:
                            imp_deb = signed[0]
                        if len(signed) > 1:
                            imp_cre = signed[1]
                    if imp_deb is None and imp_cre is None:
                        continue
                    rows.append({'Banca': banca, 'File': filename, 'Pagina': page_idx, 'Data versamento': data_vers, 'Sezione': sezione, 'Codice tributo': trib, 'Descrizione': desc, 'Anno': anno, 'Importo Debito (EUR)': imp_deb, 'Importo Credito (EUR)': imp_cre})
            except (ValueError, TypeError, AttributeError, KeyError) as e:
                logger.warning('Pagina %d saltata per errore atteso: %s', page_idx, type(e).__name__)
            except Exception as e:
                logger.error('Pagina %d - errore inatteso: %s', page_idx, type(e).__name__)
            finally:
                stampati = _totali_stampati(page_text)
                if stampati:
                    controllate += 1
                    t = totali(rows[inizio:])
                    if (t['debiti'], t['crediti'], t['saldo']) != stampati:
                        non_tornano.append(page_idx)
                page.close()
        logger.info('Parsing completato: %d righe estratte', len(rows))
    if stats is not None:
        stats['pages'] = total
        stats['text_pages'] = text_pages
        stats['deleghe_controllate'] = controllate
        stats['pagine_non_tornano'] = non_tornano
    gc.collect()
    return rows
_FILL_MAP = {k: PatternFill('solid', fgColor=v.lstrip('#')) for k, v in SEZIONE_COLORI.items()}
_HDR_FONT = Font(bold=True, color='FFFFFF')
_HDR_FILL = PatternFill('solid', fgColor='1F3864')
_HDR_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
_THIN = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
_EURO_FMT = '#,##0.00 €;-#,##0.00 €'
_TOT_BORDER = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='medium'), bottom=Side(style='thin'))

def totali(rows):
    debiti = sum((r['Importo Debito (EUR)'] or 0 for r in rows))
    crediti = sum((r['Importo Credito (EUR)'] or 0 for r in rows))
    return {'debiti': round(debiti, 2), 'crediti': round(crediti, 2), 'saldo': round(debiti - crediti, 2)}

def periodo_versamenti(rows):
    date_ = sorted((date(*map(int, reversed(d.split('/')))) for d in (r.get('Data versamento') for r in rows) if d))
    if not date_:
        return (None, None)
    return (date_[0].strftime('%d/%m/%Y'), date_[-1].strftime('%d/%m/%Y'))

def _riga_totale(ws, tot, credito_negativo):
    intest = [c.value for c in ws[1]]
    r = ws.max_row + 2
    credito = -tot['crediti'] if credito_negativo and tot['crediti'] else tot['crediti']
    ws.cell(r, 1, 'TOTALE')
    for col, v in (('Debito', tot['debiti']), ('Credito', credito), ('Netto (EUR)', tot['saldo'])):
        ws.cell(r, intest.index(col) + 1, float(v)).number_format = _EURO_FMT
    for cella in ws[r]:
        cella.font = Font(bold=True)
        cella.border = _TOT_BORDER

def _style_sheet(ws, section_col_idx=None):
    for cell in ws[1]:
        cell.font = _HDR_FONT
        cell.fill = _HDR_FILL
        cell.alignment = _HDR_ALIGN
        cell.border = _THIN
    ws.row_dimensions[1].height = 28
    for row in ws.iter_rows(min_row=2):
        fill = None
        if section_col_idx is not None:
            sezione = str(row[section_col_idx].value or '?')
            fill = _FILL_MAP.get(sezione, _FILL_MAP['?'])
        for cell in row:
            if fill:
                cell.fill = fill
            cell.border = _THIN
            cell.alignment = Alignment(vertical='center')
            if isinstance(cell.value, float):
                cell.number_format = _EURO_FMT
            elif isinstance(cell.value, date):
                cell.number_format = 'DD/MM/YYYY'
    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 55)
    ws.freeze_panes = 'A2'

def _build_dataframes(rows, credito_negativo=False):
    df = pd.DataFrame(rows)
    if 'Data versamento' in df.columns:
        df['Data versamento'] = [pd.to_datetime(v, format='%d/%m/%Y').date() if v else None for v in df['Data versamento']]
    agg = df.groupby(['Banca', 'Sezione', 'Codice tributo', 'Descrizione'], dropna=False).agg(Debito=('Importo Debito (EUR)', 'sum'), Credito=('Importo Credito (EUR)', 'sum')).reset_index()
    agg_anno = df.groupby(['Banca', 'Anno', 'Sezione'], dropna=False).agg(Debito=('Importo Debito (EUR)', 'sum'), Credito=('Importo Credito (EUR)', 'sum')).reset_index()
    agg['Netto (EUR)'] = agg['Debito'].fillna(0) - agg['Credito'].fillna(0)
    agg_anno['Netto (EUR)'] = agg_anno['Debito'].fillna(0) - agg_anno['Credito'].fillna(0)
    for t in (agg, agg_anno):
        for col in ('Debito', 'Credito', 'Netto (EUR)'):
            t[col] = t[col].map(lambda v: round(v, 2) if isinstance(v, float) else v)
    if credito_negativo and 'Importo Credito (EUR)' in df.columns:
        neg = lambda x: -x if pd.notna(x) and x != 0 else x
        df['Importo Credito (EUR)'] = df['Importo Credito (EUR)'].apply(neg)
        agg['Credito'] = agg['Credito'].apply(neg)
        agg_anno['Credito'] = agg_anno['Credito'].apply(neg)
    return (df, agg, agg_anno)

def _write_workbook(writer, df, agg, agg_anno, tot, credito_negativo):
    det_cols = df.columns.tolist()
    sec_idx = det_cols.index('Sezione') if 'Sezione' in det_cols else None
    df.to_excel(writer, index=False, sheet_name='Dettaglio')
    agg.to_excel(writer, index=False, sheet_name='Riepilogo')
    agg_anno.to_excel(writer, index=False, sheet_name='Per Anno')
    wb = writer.book
    _style_sheet(wb['Dettaglio'], section_col_idx=sec_idx)
    _style_sheet(wb['Riepilogo'], section_col_idx=1)
    _style_sheet(wb['Per Anno'], section_col_idx=2)
    _riga_totale(wb['Riepilogo'], tot, credito_negativo)
    _riga_totale(wb['Per Anno'], tot, credito_negativo)
    _firma_datace(wb['Dettaglio'])

def export_excel(rows, out_path, credito_negativo=False):
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    df, agg, agg_anno = _build_dataframes(rows, credito_negativo)
    with pd.ExcelWriter(out_path, engine='openpyxl') as xw:
        _write_workbook(xw, df, agg, agg_anno, totali(rows), credito_negativo)
    return out_path

def export_excel_bytes(rows, credito_negativo=False) -> bytes:
    df, agg, agg_anno = _build_dataframes(rows, credito_negativo)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as xw:
        _write_workbook(xw, df, agg, agg_anno, totali(rows), credito_negativo)
    return buf.getvalue()
