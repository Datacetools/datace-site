import base64
import io
import math
import re
import logging
import xml.etree.ElementTree as ET
from datetime import date
from itertools import islice
from pathlib import Path
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.styles.cell_style import StyleArray
from openpyxl.utils import get_column_letter
logger = logging.getLogger('FatturaXML')
FIRMA_TESTO = 'Datace · datacetools.com'
FIRMA_LINK = 'https://datacetools.com/?utm_source=excel&utm_medium=file'
FIRMA_LOGO = Path(__file__).parent / 'firma-datace.png'

def _firma_datace(ws):
    try:
        from openpyxl.drawing.image import Image
        from openpyxl.worksheet.cell_range import CellRange
        from openpyxl.utils.cell import coordinate_from_string
        logo = Image(str(FIRMA_LOGO))
        logo.width = logo.height = 28
        unite = [str(r) for r in ws.merged_cells.ranges]
        for r in unite:
            ws.unmerge_cells(r)
        altezze = {i: d.height for i, d in ws.row_dimensions.items() if d.height}
        ws.insert_rows(1)
        for r in unite:
            cr = CellRange(r)
            cr.shift(row_shift=1)
            ws.merge_cells(cr.coord)
        for i in sorted(altezze, reverse=True):
            ws.row_dimensions[i + 1].height = altezze[i]
            ws.row_dimensions[i].height = None
        if ws.freeze_panes:
            col, riga = coordinate_from_string(ws.freeze_panes)
            ws.freeze_panes = f'{col}{riga + 1}'
        ws.row_dimensions[1].height = 30
        cella = ws.cell(1, 1, FIRMA_TESTO)
        cella.hyperlink = FIRMA_LINK
        cella.font = Font(size=11, bold=True, color='FFFFFF', underline='single')
        cella.fill = PatternFill('solid', fgColor='0F2744')
        cella.alignment = Alignment(indent=6, vertical='center')
        if ws.max_column > 1:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ws.max_column)
        from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
        logo.anchor = TwoCellAnchor(editAs='twoCell', _from=AnchorMarker(col=0, colOff=45000, row=0, rowOff=45000), to=AnchorMarker(col=0, colOff=335000, row=0, rowOff=335000))
        ws.add_image(logo)
    except Exception as e:
        logger.warning('Banner non inserito: %s', type(e).__name__)
TIPO_DOCUMENTO = {'TD01': 'Fattura', 'TD02': 'Acconto/anticipo su fattura', 'TD03': 'Acconto/anticipo su parcella', 'TD04': 'Nota di credito', 'TD05': 'Nota di debito', 'TD06': 'Parcella', 'TD16': 'Integrazione reverse charge interno', 'TD17': 'Integrazione/autofattura acquisto servizi estero', 'TD18': 'Integrazione acquisto beni intracomunitari', 'TD19': 'Integrazione/autofattura acquisto beni ex art.17 c.2', 'TD20': 'Autofattura per regolarizzazione', 'TD21': 'Autofattura per splafonamento', 'TD22': 'Estrazione beni da deposito IVA', 'TD23': 'Estrazione beni da deposito IVA con versamento', 'TD24': 'Fattura differita (art.21 c.4 lett.a)', 'TD25': 'Fattura differita (art.21 c.4 terzo periodo lett.b)', 'TD26': 'Cessione di beni ammortizzabili', 'TD27': 'Fattura per autoconsumo o cessioni gratuite', 'TD28': 'Acquisti da San Marino con IVA (fattura cartacea)', 'TD29': 'Comunicazione omessa o irregolare fatturazione (art.6 c.8 D.Lgs.471/97)'}
MODALITA_PAGAMENTO = {'MP01': 'Contanti', 'MP02': 'Assegno', 'MP03': 'Assegno circolare', 'MP04': 'Contanti presso Tesoreria', 'MP05': 'Bonifico', 'MP06': 'Vaglia cambiario', 'MP07': 'Bollettino bancario', 'MP08': 'Carta di pagamento', 'MP09': 'RID', 'MP10': 'RID utenze', 'MP11': 'RID veloce', 'MP12': 'RIBA', 'MP13': 'MAV', 'MP14': 'Quietanza erario', 'MP15': 'Giroconto conti speciali', 'MP16': 'Domiciliazione bancaria', 'MP17': 'Domiciliazione postale', 'MP18': 'Bollettino di c/c postale', 'MP19': 'SEPA Direct Debit', 'MP20': 'SEPA Direct Debit CORE', 'MP21': 'SEPA Direct Debit B2B', 'MP22': 'Trattenuta su somme già riscosse', 'MP23': 'PagoPA'}
NATURA = {'N1': 'Escluse ex art.15', 'N2': 'Non soggette', 'N2.1': 'Non soggette artt. 7-7septies', 'N2.2': 'Non soggette - altri casi', 'N3': 'Non imponibili', 'N3.1': 'Non imponibili - esportazioni', 'N3.2': 'Non imponibili - cessioni intracomunitarie', 'N3.3': 'Non imponibili - cessioni San Marino', 'N3.4': 'Non imponibili - operazioni assimilate esportazioni', 'N3.5': 'Non imponibili - dichiarazioni intento', 'N3.6': 'Non imponibili - altre operazioni', 'N4': 'Esenti', 'N5': 'Regime del margine', 'N6': 'Inversione contabile (reverse charge)', 'N6.1': 'RC - cessione rottami', 'N6.2': 'RC - oro e argento', 'N6.3': 'RC - subappalto edilizia', 'N6.4': 'RC - cessione fabbricati', 'N6.5': 'RC - telefoni cellulari', 'N6.6': 'RC - prodotti elettronici', 'N6.7': 'RC - comparto edile', 'N6.8': 'RC - settore energetico', 'N6.9': 'RC - altri casi', 'N7': 'IVA assolta in altro Stato UE'}
ESIGIBILITA = {'I': 'Immediata', 'D': 'Differita', 'S': 'Scissione pagamenti (split payment)'}
TIPO_NOTIFICA = {'RicevutaConsegna': 'RC - Ricevuta Consegna', 'NotificaScarto': 'NS - Notifica Scarto', 'AttestatoMancataConsegna': 'MC - Mancata Consegna', 'NotificaMancataConsegna': 'MC - Mancata Consegna', 'EsitoCommittente': 'EC - Esito Committente', 'DecorrenzaTermini': 'DT - Decorrenza Termini', 'NotificaEsitoCommittente': 'EC - Esito Committente', 'NotificaEsito': 'EC - Esito Committente', 'ScartoEsitoCommittente': 'SE - Scarto Esito Committente', 'NotificaDecorrenzaTermini': 'DT - Decorrenza Termini', 'AttestazioneTrasmissioneFattura': 'AT - Attestazione Trasmissione', 'RicevutaScarto': 'NS - Ricevuta Scarto', 'RicevutaImpossibilitaRecapito': 'MC - Impossibilità di Recapito'}
NON_LETTI = {'FatturaElettronicaSemplificata': ('fattura semplificata', 'semplificate'), 'DatiCorrispettivi': ('corrispettivi', 'corrispettivi')}

class FormatoNonSupportato(ValueError):

    def __init__(self, radice: str):
        nome, self.cartella = NON_LETTI[radice]
        super().__init__(f'{nome}: formato non supportato da questo strumento')
ESITO_EC = {'EC01': 'Accettata dal committente', 'EC02': 'Rifiutata dal committente', 'EN00': 'Notifica non conforme al formato', 'EN01': 'Notifica non ammissibile'}

def _lname(el) -> str:
    return el.tag.rsplit('}', 1)[-1]

def _first(el, name):
    if el is None:
        return None
    for c in el:
        if _lname(c) == name:
            return c
    return None

def _all(el, name):
    if el is None:
        return []
    return [c for c in el if _lname(c) == name]

def _path(el, *names):
    cur = el
    for n in names:
        cur = _first(cur, n)
        if cur is None:
            return None
    return cur

def _text(el, *names) -> str:
    node = _path(el, *names) if names else el
    if node is None or node.text is None:
        return ''
    return node.text.strip()

def _num(s):
    if not s:
        return None
    s = s.strip()
    try:
        v = float(s)
    except ValueError:
        try:
            v = float(s.replace('.', '').replace(',', '.'))
        except ValueError:
            return None
    return v if math.isfinite(v) else None

def _date(s: str) -> str:
    if not s:
        return ''
    return s.split('T')[0] if 'T' in s else s

def _clean_xml_bytes(data: bytes) -> bytes:
    enc_match = re.search(b'encoding\\s*=\\s*[\'"]([^\'"]+)[\'"]', data[:200])
    declared_enc = enc_match.group(1).decode('ascii', errors='ignore').upper() if enc_match else 'UTF-8'
    if declared_enc in ('UTF-8', 'UTF8'):
        text = data.decode('utf-8', errors='replace')
    else:
        try:
            text = data.decode(declared_enc, errors='replace')
        except (LookupError, UnicodeDecodeError):
            text = data.decode('latin-1', errors='replace')
        text = re.sub('(encoding\\s*=\\s*[\'"])[^\'"]+([\'"])', '\\g<1>UTF-8\\g<2>', text, count=1)
    text = re.sub('[\\x00-\\x08\\x0b\\x0c\\x0e-\\x1f]', '', text)
    return text.encode('utf-8')

def _tlv(data: bytes, pos: int):
    tag, n = (data[pos], data[pos + 1])
    pos += 2
    if n == 128:
        return (tag, None, pos)
    if n & 128:
        k = n & 127
        n = int.from_bytes(data[pos:pos + k], 'big')
        pos += k
    return (tag, n, pos)

def _fine(data: bytes, pos: int) -> int:
    _, n, p = _tlv(data, pos)
    if n is not None:
        return p + n
    while data[p:p + 2] != b'\x00\x00':
        p = _fine(data, p)
    return p + 2

def _figli(data: bytes, pos: int):
    _, n, p = _tlv(data, pos)
    fine = len(data) if n is None else p + n
    while p < fine and data[p:p + 2] != b'\x00\x00':
        yield p
        p = _fine(data, p)

def _octet(data: bytes, pos: int) -> bytes:
    tag, n, p = _tlv(data, pos)
    if tag & 31 != 4:
        raise ValueError('Atteso un OCTET STRING nella busta .p7m.')
    if tag & 32:
        return b''.join((_octet(data, q) for q in _figli(data, pos)))
    return data[p:p + n]

def _contenuto_p7m(data: bytes) -> bytes:
    figlio = lambda pos, i: next(islice(_figli(data, pos), i, None))
    firmato = figlio(figlio(0, 1), 0)
    incapsulato = figlio(firmato, 2)
    return _octet(data, figlio(figlio(incapsulato, 1), 0))

def _xml_firmato(data: bytes) -> bytes:
    if data[:1] != b'0':
        try:
            corpo = b''.join((r for r in data.splitlines() if not r.startswith(b'-----')))
            decodificato = base64.b64decode(corpo)
            if decodificato[:1] == b'0':
                data = decodificato
        except ValueError:
            pass
    try:
        contenuto = _contenuto_p7m(data)
    except Exception:
        contenuto = _xml_cercato_nei_byte(data)
    return contenuto
_BOM_E_SPAZI = b'\xef\xbb\xbf \t\r\n'

def _xml_originale(data: bytes) -> bytes:
    for _ in range(3):
        if data.lstrip(_BOM_E_SPAZI).startswith(b'<'):
            return data
        data = _xml_firmato(data)
    return data
_VIETATI_NEI_NOMI = re.compile('[\\x00-\\x1f<>:"|?*]')

def _nome_sicuro(nome: str, ripiego: str) -> str:
    nome = nome.replace('\\', '/').rsplit('/', 1)[-1]
    return _VIETATI_NEI_NOMI.sub('_', nome).strip(' .')[:120] or ripiego

def estrai_file(data: bytes, filename: str) -> list:
    base = _nome_sicuro(re.sub('(?i)(\\.xml)?(\\.p7m)*$', '', filename), 'fattura')
    xml = _xml_originale(data)
    file = [(f'xml/{base}.xml', xml)]
    root = ET.fromstring(_clean_xml_bytes(xml.lstrip(_BOM_E_SPAZI)))
    allegati = (a for body in _all(root, 'FatturaElettronicaBody') for a in _all(body, 'Allegati'))
    for n, al in enumerate(allegati, start=1):
        try:
            contenuto = base64.b64decode(_text(al, 'Attachment'))
        except ValueError:
            logger.warning('Allegato %d non decodificabile: saltato', n)
            continue
        if not contenuto:
            continue
        nome = _nome_sicuro(_text(al, 'NomeAttachment'), f'allegato_{n}')
        formato = _text(al, 'FormatoAttachment').lower()
        if '.' not in nome and formato.isalnum():
            nome += f'.{formato}'
        file.append((f'allegati/{base}/{nome}', contenuto))
    return file

def _xml_cercato_nei_byte(data: bytes) -> bytes:
    start = data.find(b'<?xml')
    if start == -1:
        start = data.find(b'<p:FatturaElettronica')
        if start == -1:
            start = data.find(b'<FatturaElettronica')
        if start == -1:
            raise ValueError('Contenuto XML non trovato nel file .p7m.')
    end = -1
    for marker in (b'</p:FatturaElettronica>', b'</FatturaElettronica>', b'FatturaElettronica>'):
        pos = data.rfind(marker)
        if pos != -1:
            end = pos + len(marker)
            break
    if end == -1:
        raise ValueError('Tag di chiusura FatturaElettronica non trovato nel .p7m.')
    return data[start:end]

def _denominazione(anagrafica) -> str:
    den = _text(anagrafica, 'Denominazione')
    if den:
        return den
    return ' '.join((p for p in (_text(anagrafica, 'Nome'), _text(anagrafica, 'Cognome')) if p))

def _soggetto(blocco) -> dict:
    dati = _first(blocco, 'DatiAnagrafici')
    anagrafica = _first(dati, 'Anagrafica')
    nome = _denominazione(anagrafica) if anagrafica is not None else ''
    piva = ''
    idfisc = _first(dati, 'IdFiscaleIVA')
    if idfisc is not None:
        paese = _text(idfisc, 'IdPaese')
        codice = _text(idfisc, 'IdCodice')
        piva = f'{paese}{codice}'.strip()
    cf = _text(dati, 'CodiceFiscale')
    return {'nome': nome, 'piva': piva, 'cf': cf}

def _parse_notifica_sdi(root, fname: str) -> list:
    root_name = _lname(root)
    tipo_desc = TIPO_NOTIFICA.get(root_name, f'Notifica SdI ({root_name})')
    notifica = {'file': fname, 'tipo': tipo_desc, 'id_sdi': _text(root, 'IdentificativoSdI'), 'nome_file_riferito': _text(root, 'NomeFile'), 'data_ricezione': _date(_text(root, 'DataOraRicezione') or _text(root, 'DataIoraRicezione')), 'data_consegna': _date(_text(root, 'DataOraConsegna') or _text(root, 'DataMessaADisposizione') or _text(root, 'DataOraDecorrenzaTermini') or _text(root, 'DataOraAttestatoDi')), 'esito': '', 'descrizione': _text(root, 'Descrizione'), 'errori': ''}
    for blocco in (root, *_all(root, 'EsitoCommittente')):
        esito_cod = _text(blocco, 'Esito') or _text(blocco, 'Scarto')
        if esito_cod:
            notifica['esito'] = ESITO_EC.get(esito_cod, esito_cod)
        notifica['id_sdi'] = notifica['id_sdi'] or _text(blocco, 'IdentificativoSdI')
        notifica['descrizione'] = notifica['descrizione'] or _text(blocco, 'Descrizione')
        rif = _first(blocco, 'RiferimentoFattura')
        num = _text(rif, 'NumeroFattura')
        if num and (not notifica['nome_file_riferito']):
            notifica['nome_file_riferito'] = f"N.{num}/{_text(rif, 'AnnoFattura')}"
    errori = []
    for lista in _all(root, 'ListaErrori'):
        for err in _all(lista, 'Errore'):
            cod = _text(err, 'Codice')
            desc = _text(err, 'Descrizione')
            errori.append(f'{cod}: {desc}' if cod else desc)
    if errori:
        notifica['errori'] = ' | '.join(errori)
    logger.info('Notifica SdI %s', tipo_desc)
    return [notifica]

def _parse_metadati_sdi(data: bytes, fname: str) -> list:
    root = ET.fromstring(data)
    meta = {}
    for m in root:
        nome_el = _first(m, 'nome')
        valore_el = _first(m, 'valore')
        if nome_el is not None and valore_el is not None:
            k = (nome_el.text or '').strip()
            v = (valore_el.text or '').strip()
            if k:
                meta[k] = v
    data_raw = meta.get('dataemissioneminima', '')
    return [{'file': fname, 'tipo_cod': 'SDI-META', 'tipo': 'Metadato SdI', 'numero': meta.get('idregistro', ''), 'data': _date(data_raw), 'divisa': 'EUR', 'totale_doc': None, 'causale': f"File SdI: {meta.get('nomefile', '')}", 'cup': '', 'cig': '', 'ordine': '', 'ufficio': '', 'fornitore': {'nome': meta.get('cedentedenominazione', ''), 'piva': meta.get('cedentepiva', ''), 'cf': meta.get('cedentecf', '')}, 'cliente': {'nome': meta.get('cessionariodenominazione', ''), 'piva': meta.get('cessionariopiva', ''), 'cf': meta.get('cessionariocf', '')}, 'righe': [], 'riepilogo_iva': [], 'pagamenti': [], 'allegati': [], 'tot_imponibile': 0.0, 'tot_imposta': 0.0, 'scadenza': '', '_notifiche': []}]

def parse_fattura(source, filename: str=None) -> list:
    if isinstance(source, (bytes, bytearray)):
        data = bytes(source)
        fname = filename or 'fattura.xml'
    elif isinstance(source, io.IOBase):
        data = source.read()
        fname = filename or 'fattura.xml'
    else:
        p = Path(source)
        data = p.read_bytes()
        fname = filename or p.name
    data = _clean_xml_bytes(_xml_originale(data).lstrip(_BOM_E_SPAZI))
    root = ET.fromstring(data)
    root_name = _lname(root)
    if root_name == 'metadatiFattura':
        logger.info('Metadati SdI')
        return _parse_metadati_sdi(data, fname)
    if root_name in NON_LETTI:
        raise FormatoNonSupportato(root_name)
    if root_name in TIPO_NOTIFICA:
        notifiche = _parse_notifica_sdi(root, fname)
        return [{'file': fname, 'tipo_cod': 'NOTIFICA', 'tipo': n['tipo'], 'numero': n['id_sdi'], 'data': n['data_ricezione'], 'fornitore': {'nome': '', 'piva': '', 'cf': ''}, 'cliente': {'nome': '', 'piva': '', 'cf': ''}, 'totale_doc': None, 'tot_imponibile': 0.0, 'tot_imposta': 0.0, 'divisa': 'EUR', 'causale': '', 'cup': '', 'cig': '', 'ordine': '', 'ufficio': '', 'righe': [], 'riepilogo_iva': [], 'pagamenti': [], 'allegati': [], 'scadenza': '', '_notifiche': notifiche} for n in notifiche]
    header = _first(root, 'FatturaElettronicaHeader')
    trasmissione = _first(header, 'DatiTrasmissione')
    cedente = _first(header, 'CedentePrestatore')
    cessionario = _first(header, 'CessionarioCommittente')
    fornitore = _soggetto(cedente) if cedente is not None else {'nome': '', 'piva': '', 'cf': ''}
    cliente = _soggetto(cessionario) if cessionario is not None else {'nome': '', 'piva': '', 'cf': ''}
    formato = _text(trasmissione, 'FormatoTrasmissione')
    ufficio = _text(trasmissione, 'CodiceDestinatario') if formato.startswith('FPA') else ''
    fatture = []
    for body in _all(root, 'FatturaElettronicaBody'):
        dati_gen_doc = _path(body, 'DatiGenerali', 'DatiGeneraliDocumento')
        dati_generali = _first(body, 'DatiGenerali')
        tipo_cod = _text(dati_gen_doc, 'TipoDocumento')
        causali = ' | '.join((c.text.strip() for c in _all(dati_gen_doc, 'Causale') if c.text))
        cup, cig, ordine = ('', '', '')
        for tag in ('DatiOrdineAcquisto', 'DatiContratto', 'DatiConvenzione', 'DatiRicezione', 'DatiFattureCollegate'):
            for item in _all(dati_generali, tag):
                cup = cup or _text(item, 'CodiceCUP')
                cig = cig or _text(item, 'CodiceCIG')
                ordine = ordine or _text(item, 'IdDocumento')
        doc = {'file': fname, 'tipo_cod': tipo_cod, 'tipo': TIPO_DOCUMENTO.get(tipo_cod, tipo_cod), 'numero': _text(dati_gen_doc, 'Numero'), 'data': _text(dati_gen_doc, 'Data'), 'divisa': _text(dati_gen_doc, 'Divisa') or 'EUR', 'totale_doc': _num(_text(dati_gen_doc, 'ImportoTotaleDocumento')), 'causale': causali, 'cup': cup, 'cig': cig, 'ordine': ordine, 'ufficio': ufficio, 'fornitore': fornitore, 'cliente': cliente, 'righe': [], 'riepilogo_iva': [], 'pagamenti': [], 'allegati': [], '_notifiche': []}
        beni = _first(body, 'DatiBeniServizi')
        for r in _all(beni, 'DettaglioLinee'):
            doc['righe'].append({'linea': _text(r, 'NumeroLinea'), 'descrizione': _text(r, 'Descrizione'), 'quantita': _num(_text(r, 'Quantita')), 'unita': _text(r, 'UnitaMisura'), 'prezzo_unitario': _num(_text(r, 'PrezzoUnitario')), 'prezzo_totale': _num(_text(r, 'PrezzoTotale')), 'aliquota': _num(_text(r, 'AliquotaIVA')), 'natura_cod': _text(r, 'Natura')})
        for rp in _all(beni, 'DatiRiepilogo'):
            nat = _text(rp, 'Natura')
            doc['riepilogo_iva'].append({'aliquota': _num(_text(rp, 'AliquotaIVA')), 'imponibile': _num(_text(rp, 'ImponibileImporto')), 'imposta': _num(_text(rp, 'Imposta')), 'natura_cod': nat, 'natura': NATURA.get(nat, nat), 'esigibilita': ESIGIBILITA.get(_text(rp, 'EsigibilitaIVA'), _text(rp, 'EsigibilitaIVA'))})
        pag = _first(body, 'DatiPagamento')
        for dp in _all(pag, 'DettaglioPagamento'):
            mp = _text(dp, 'ModalitaPagamento')
            doc['pagamenti'].append({'modalita_cod': mp, 'modalita': MODALITA_PAGAMENTO.get(mp, mp), 'scadenza': _text(dp, 'DataScadenzaPagamento'), 'importo': _num(_text(dp, 'ImportoPagamento'))})
        for al in _all(body, 'Allegati'):
            doc['allegati'].append({'nome': _text(al, 'NomeAttachment'), 'formato': _text(al, 'FormatoAttachment'), 'descrizione': _text(al, 'DescrizioneAttachment'), 'compressione': _text(al, 'AlgoritmoCompressione')})
        doc['tot_imponibile'] = sum((x['imponibile'] or 0 for x in doc['riepilogo_iva']))
        doc['tot_imposta'] = sum((x['imposta'] or 0 for x in doc['riepilogo_iva']))
        expected_total = round(doc['tot_imponibile'] + doc['tot_imposta'], 2)
        if doc['totale_doc'] is None:
            doc['totale_doc'] = expected_total
        elif abs(doc['totale_doc'] - expected_total) > 0.02:
            logger.warning('Totale documento non quadra con la somma delle righe: usato quello dichiarato')
        scad = [p['scadenza'] for p in doc['pagamenti'] if p.get('scadenza')]
        doc['scadenza'] = max(scad) if scad else ''
        fatture.append(doc)
    logger.info('Parsing %s: %d fattura/e', root_name if fatture else 'XML non riconosciuto', len(fatture))
    return fatture

def _build_dataframes(fatture: list):
    riepilogo_rows = []
    righe_rows = []
    iva_rows = []
    notifiche_rows = []
    allegati_rows = []
    for f in fatture:
        forn = f['fornitore']
        cli = f['cliente']
        tipo_cod = f.get('tipo_cod', '')
        if tipo_cod == 'NOTIFICA':
            for n in f.get('_notifiche', []):
                notifiche_rows.append({'File': n['file'], 'Tipo notifica': n['tipo'], 'ID SdI': n['id_sdi'], 'File fattura riferita': n['nome_file_riferito'], 'Data ricezione': n['data_ricezione'], 'Data consegna/scarto': n['data_consegna'], 'Esito': n['esito'], 'Descrizione': n['descrizione'], 'Errori': n['errori']})
            continue
        if tipo_cod == 'SDI-META':
            riepilogo_rows.append({'File': f['file'], 'Tipo documento': 'Metadato SdI', 'Numero/Riferimento': f['numero'], 'Data': f['data'], 'Fornitore': forn['nome'], 'P.IVA/CF Fornitore': forn['piva'] or forn['cf'], 'Cliente': cli['nome'], 'P.IVA/CF Cliente': cli['piva'] or cli['cf'], 'Imponibile': '', 'IVA': '', 'Totale': '', 'Divisa': 'EUR', 'Scadenza': '', 'CUP': '', 'CIG': '', 'N. Ordine/Contratto': '', 'Ufficio destinatario': '', 'Note': f.get('causale', '')})
            continue
        riepilogo_rows.append({'File': f['file'], 'Tipo documento': f['tipo'], 'Numero/Riferimento': f['numero'], 'Data': f['data'], 'Fornitore': forn['nome'], 'P.IVA/CF Fornitore': forn['piva'] or forn['cf'], 'Cliente': cli['nome'], 'P.IVA/CF Cliente': cli['piva'] or cli['cf'], 'Imponibile': round(f['tot_imponibile'], 2) if f['tot_imponibile'] is not None else None, 'IVA': round(f['tot_imposta'], 2) if f['tot_imposta'] is not None else None, 'Totale': f['totale_doc'], 'Divisa': f['divisa'], '_divisa': f['divisa'], 'Scadenza': f['scadenza'], 'CUP': f.get('cup', ''), 'CIG': f.get('cig', ''), 'N. Ordine/Contratto': f.get('ordine', ''), 'Ufficio destinatario': f.get('ufficio', ''), 'Note': f.get('causale', '')})
        for r in f['righe']:
            righe_rows.append({'File': f['file'], 'Numero fattura': f['numero'], 'Data': f['data'], 'Fornitore': forn['nome'], 'Linea': r['linea'], 'Descrizione': r['descrizione'], 'Quantita': r['quantita'], 'U.M.': r['unita'], 'Prezzo unitario': r['prezzo_unitario'], 'Prezzo totale': r['prezzo_totale'], 'Aliquota IVA %': r['aliquota'], 'Natura': NATURA.get(r['natura_cod'], r['natura_cod']), '_divisa': f['divisa']})
        for v in f['riepilogo_iva']:
            iva_rows.append({'File': f['file'], 'Numero fattura': f['numero'], 'Data': f['data'], 'Aliquota IVA %': v['aliquota'], 'Imponibile': v['imponibile'], 'Imposta': v['imposta'], 'Natura': v['natura'], 'Esigibilita IVA': v['esigibilita'], '_divisa': f['divisa']})
        for al in f.get('allegati', []):
            allegati_rows.append({'File': f['file'], 'Numero fattura': f['numero'], 'Data': f['data'], 'Nome allegato': al['nome'], 'Formato': al['formato'], 'Descrizione': al['descrizione'], 'Compressione': al['compressione']})
    return (pd.DataFrame(riepilogo_rows), pd.DataFrame(righe_rows), pd.DataFrame(iva_rows), pd.DataFrame(notifiche_rows), pd.DataFrame(allegati_rows))
_COLORS = {'navy': '1F3864', 'teal': '1A6B7E', 'purple': '4B2E7E', 'green': '1A5C3A', 'orange': '7E3B1A'}
_HDR_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
_THIN = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
_EURO_FMT = '#,##0.00 €'
_NUM_FMT = '#,##0.00'
_DATE_FMT = 'DD/MM/YYYY'
_IMPORTI = {'Imponibile', 'IVA', 'Totale', 'Prezzo unitario', 'Prezzo totale', 'Imposta'}

def _style_sheet(ws, hdr_color: str='navy', divise: list | None=None):
    if ws.max_row < 1:
        return
    hdr_fill = PatternFill('solid', fgColor=_COLORS.get(hdr_color, '1F3864'))
    hdr_font = Font(bold=True, color='FFFFFF')
    for cell in ws[1]:
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = _HDR_ALIGN
        cell.border = _THIN
    ws.row_dimensions[1].height = 30
    stile = None
    titoli = [c.value or '' for c in ws[1]]
    for i, row in enumerate(ws.iter_rows(min_row=2)):
        in_euro = divise is None or divise[i] == 'EUR'
        if stile is None:
            stile = (ws.parent._borders.add(_THIN), ws.parent._alignments.add(Alignment(vertical='center')))
        for cell in row:
            if not cell._style:
                cell._style = StyleArray()
            cell._style.borderId, cell._style.alignmentId = stile
            hdr = titoli[cell.column - 1]
            if cell.data_type == 'f':
                cell.data_type = 's'
            if isinstance(cell.value, float):
                if hdr in _IMPORTI:
                    cell.number_format = _EURO_FMT if in_euro else _NUM_FMT
                elif '%' in str(hdr):
                    cell.number_format = '0.00'
            elif isinstance(cell.value, str) and (str(hdr).startswith('Data') or hdr == 'Scadenza'):
                try:
                    cell.value = date.fromisoformat(cell.value[:10])
                    cell.number_format = _DATE_FMT
                except ValueError:
                    pass
    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 55)
    ws.freeze_panes = 'A2'

def _write_workbook(writer, df_riepilogo, df_righe, df_iva, df_notifiche, df_allegati):
    sheets = [('Riepilogo Fatture', df_riepilogo, 'navy'), ('Dettaglio Righe', df_righe, 'teal'), ('Riepilogo IVA', df_iva, 'purple'), ('Notifiche SdI', df_notifiche, 'green'), ('Allegati', df_allegati, 'orange')]
    for name, df, color in sheets:
        divise = df.pop('_divisa').tolist() if '_divisa' in df else None
        df.to_excel(writer, index=False, sheet_name=name)
        _style_sheet(writer.book[name], color, divise)
    _firma_datace(writer.book.worksheets[0])

def export_excel(fatture: list, out_path):
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    dfs = _build_dataframes(fatture)
    with pd.ExcelWriter(out_path, engine='openpyxl') as xw:
        _write_workbook(xw, *dfs)
    return out_path

def export_excel_bytes(fatture: list) -> bytes:
    dfs = _build_dataframes(fatture)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as xw:
        _write_workbook(xw, *dfs)
    return buf.getvalue()
